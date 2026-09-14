#!/usr/bin/env python3
"""Markdown → 美化 docx（应用主题样式）。

用法::

    python3 md2docx.py 输入.md --theme business-blue --title "标题" --out 产出物/xxx.docx
    cat 输入.md | python3 md2docx.py - --out 产出物/xxx.docx
"""
from __future__ import annotations

import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import themes as _themes   # noqa: E402
import _md                 # noqa: E402
import docx_theme          # noqa: E402


def build(md_text: str, title: str, theme: dict, out_path: str) -> str:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    docx_theme.apply_theme(doc, theme)

    if title:
        h = doc.add_heading(title, level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in h.runs:
            r.font.color.rgb = RGBColor.from_string(
                docx_theme.hex6(theme.get("color_primary")))

    for blk in _md.parse(md_text):
        kind = blk[0]
        if kind == "h":
            doc.add_heading(blk[2], level=min(blk[1], 4))
        elif kind == "p":
            doc.add_paragraph(blk[1])
        elif kind == "ul":
            doc.add_paragraph(blk[1], style="List Bullet")
        elif kind == "ol":
            doc.add_paragraph(blk[1], style="List Number")
        elif kind == "code":
            p = doc.add_paragraph()
            r = p.add_run(blk[2])
            r.font.name = theme.get("font_mono", "Courier New")
            r.font.size = Pt(9)
        elif kind == "img":
            try:
                doc.add_picture(blk[2])
            except Exception:
                doc.add_paragraph("[图片未找到: %s]" % blk[2])
        elif kind == "table":
            rows = blk[1]
            cols = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=cols)
            t.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci in range(cols):
                    t.rows[ri].cells[ci].text = row[ci] if ci < len(row) else ""
            docx_theme.style_table(t, theme)

    d = os.path.dirname(os.path.abspath(out_path))
    if d:
        os.makedirs(d, exist_ok=True)
    doc.save(out_path)
    return out_path


def main() -> int:
    ap = argparse.ArgumentParser(description="Markdown → 美化 docx")
    ap.add_argument("input", help="Markdown 文件（- 表示 stdin）")
    ap.add_argument("--theme", default="business-blue", help="主题名（themes.py list）")
    ap.add_argument("--title", default="", help="文档标题")
    ap.add_argument("--out", default="", help="输出 docx 路径")
    a = ap.parse_args()

    md = sys.stdin.read() if a.input == "-" else open(a.input, encoding="utf-8").read()
    theme = _themes.load(a.theme)
    out = a.out or (os.path.splitext(os.path.basename(a.input))[0] + ".docx")
    path = build(md, a.title, theme, out)
    print("[OK] 已生成 %s（主题：%s）" % (path, theme.get("label", a.theme)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
