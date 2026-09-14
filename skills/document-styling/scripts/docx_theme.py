#!/usr/bin/env python3
"""python-docx 主题应用：样式、页边距、页脚页码、表格底纹。

对外只需两个函数：``apply_theme(doc, theme)`` 与 ``style_table(table, theme)``。
"""
from __future__ import annotations

import re

_FULL_HEX = re.compile(r"^[0-9a-fA-F]{6}$")


def hex6(color, default: str = "000000") -> str:
    """把 ``#RRGGBB`` / ``RRGGBB`` 归一化为无 # 的 6 位十六进制。"""
    if not color:
        return default
    c = str(color).strip().lstrip("#")
    return c if _FULL_HEX.match(c) else default


def _set_style_font(style, zh: str, en: str) -> None:
    """同时设置西文与中文（eastAsia）字体，保证中文字形生效。"""
    try:
        from docx.oxml.ns import qn
        style.font.name = en
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:ascii"), en)
        rfonts.set(qn("w:hAnsi"), en)
        rfonts.set(qn("w:eastAsia"), zh)
    except Exception:
        pass


def apply_theme(doc, theme: dict):
    """把主题样式应用到整个文档（字体/字号/配色/行距/边距/页脚页码）。"""
    from docx.shared import Pt, Mm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    zh = theme.get("font_zh", "宋体")
    en = theme.get("font_en", "Times New Roman")
    hzh = theme.get("font_heading_zh", zh)
    hen = theme.get("font_heading_en", en)
    primary = hex6(theme.get("color_primary"))
    body_size = float(theme.get("size_body", 12))
    line_spacing = float(theme.get("line_spacing", 1.5))
    para_space = float(theme.get("para_space", 6))

    # —— 正文 ——
    normal = doc.styles["Normal"]
    normal.font.size = Pt(body_size)
    normal.font.color.rgb = RGBColor.from_string(hex6(theme.get("color_text"), "1A1A1A"))
    _set_style_font(normal, zh, en)
    pf = normal.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_after = Pt(para_space)

    # —— 标题 ——
    sizes = {1: theme.get("size_h1", 22), 2: theme.get("size_h2", 16),
             3: theme.get("size_h3", 14), 4: theme.get("size_h3", 14),
             5: theme.get("size_h3", 14), 6: theme.get("size_h3", 14)}
    for lvl in range(1, 7):
        try:
            st = doc.styles["Heading %d" % lvl]
        except KeyError:
            continue
        st.font.size = Pt(float(sizes[lvl]))
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(primary)
        _set_style_font(st, hzh, hen)
        st.paragraph_format.space_before = Pt(10 if lvl == 1 else 8)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.3

    try:
        t = doc.styles["Title"]
        t.font.size = Pt(float(theme.get("size_h1", 22)) + 6)
        t.font.bold = True
        t.font.color.rgb = RGBColor.from_string(primary)
        _set_style_font(t, hzh, hen)
    except KeyError:
        pass

    # —— 页边距 ——
    margin = float(theme.get("margin_mm", 25.4))
    for sec in doc.sections:
        sec.top_margin = Mm(margin)
        sec.bottom_margin = Mm(margin)
        sec.left_margin = Mm(margin)
        sec.right_margin = Mm(margin)

    # —— 页脚页码（PAGE 域） ——
    for sec in doc.sections:
        footer = sec.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if not p.runs:
            run = p.add_run()
            f1 = OxmlElement("w:fldChar")
            f1.set(qn("w:fldCharType"), "begin")
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = "PAGE"
            f2 = OxmlElement("w:fldChar")
            f2.set(qn("w:fldCharType"), "end")
            run._r.append(f1)
            run._r.append(it)
            run._r.append(f2)
            run.font.size = Pt(float(theme.get("size_small", 10.5)))
            run.font.color.rgb = RGBColor.from_string(hex6(theme.get("color_muted"), "6B7280"))
    return doc


def _shade(cell, fill_hex: str) -> None:
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def style_table(table, theme: dict) -> None:
    """表格：表头底纹 + 白字加粗 + 隔行底色。"""
    from docx.shared import Pt, RGBColor
    head_fill = hex6(theme.get("color_table_head"))
    band_fill = hex6(theme.get("color_table_band"))
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            if r_idx == 0:
                _shade(cell, head_fill)
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor.from_string("FFFFFF")
            elif band_fill and r_idx % 2 == 0:
                _shade(cell, band_fill)
