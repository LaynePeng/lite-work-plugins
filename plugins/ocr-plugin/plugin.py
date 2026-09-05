"""OCR 工具：把图片 / PDF 页面 / PPT 内嵌图片中的文字提取出来。

- ocr_image   : 直接识别单张图片（png/jpg/bmp/tiff/webp 等）
- ocr_document: PDF 每页渲染成图片后逐页识别（pymupdf，无需 poppler）
- ocr_pptx    : 提取 PPT 每页内嵌图片并识别，带位置信息（left/top/宽/高）
                并按"先上后下、先左后右"排序，还原阅读顺序

OCR 引擎：rapidocr-onnxruntime（纯 pip 安装、离线可用、内置中英文模型，
无需系统级 Tesseract/poppler）。依赖缺失时返回可操作提示。
"""
# 同步自 lite-work 主仓库 litework/tools/ocr.py（社区独立分发版）
from __future__ import annotations

import logging
import os
import re
from typing import Any, Dict, List, Optional

from litework.core.types import ToolDefinition

logger = logging.getLogger("litework.tools.ocr")

# ---------------------------------------------------------------- 依赖检查

_HAS_RAPIDOCR: bool = False
_HAS_PYMUPDF: bool = False

try:
    from rapidocr_onnxruntime import RapidOCR

    _HAS_RAPIDOCR = True
except ImportError:
    pass

try:
    import pymupdf  # 新版包名（旧名 fitz 已弃用）

    _HAS_PYMUPDF = True
except ImportError:
    try:
        import fitz as pymupdf  # type: ignore[no-redef]

        _HAS_PYMUPDF = True
    except ImportError:
        pass

# 单例：RapidOCR 模型加载约数百 MB 内存，进程内复用
_OCR_ENGINE: Any = None

MAX_TEXT_LEN = 4000  # 单图识别结果截断长度（上下文保护）


def _get_engine():
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE


def _missing_ocr_msg() -> str:
    return (
        "[OCR] 需要安装 OCR 引擎才能使用此工具。\n"
        "请运行: pip install rapidocr-onnxruntime pymupdf\n"
        "（rapidocr-onnxruntime 自带中英文模型，离线可用，无需系统装 Tesseract）"
    )


def _missing_pdf_msg() -> str:
    return "[OCR] 识别 PDF 需要 pymupdf，请运行: pip install pymupdf"


# ---------------------------------------------------------------- 核心识别

def _ocr_image_bytes(image_path: str) -> str:
    """识别单张图片，返回文字（按行拼接）。"""
    engine = _get_engine()
    # RapidOCR 返回 [(box, text, score), ...]，按 top-left y 排序保持阅读顺序
    result, _elapse = engine(image_path)
    if not result:
        return ""
    lines = []
    for box, text, _score in result:
        try:
            y = float(box[0][1])
        except Exception:
            y = 0.0
        lines.append((y, text))
    lines.sort(key=lambda t: t[0])
    return "\n".join(text for _, text in lines if text and str(text).strip())


# ---------------------------------------------------------------- OCRTools

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif", ".webp", ".gif", ".ico"}


class OCRTools:
    def __init__(self, workspace: Optional[str]) -> None:
        self.workspace = os.path.abspath(workspace) if workspace else os.path.expanduser("~")

    # ------------------------------------------------------------ 工具定义

    def get_tools(self) -> List[ToolDefinition]:
        return [
            ToolDefinition(
                name="ocr_image",
                description=(
                    "识别图片中的文字（OCR）。支持 PNG/JPG/BMP/TIFF/WebP 等图片，"
                    "也支持工作区内的 .uploads/ 素材。返回提取出的文字内容。"
                ),
                parameters={
                    "type": "object",
                    "properties": {
                        "path": {
                            "type": "string",
                            "description": "图片文件路径（相对工作区）",
                        },
                    },
                    "required": ["path"],
                },
            ),
            ToolDefinition(
                name="ocr_document",
                description=(
                    "识别 PDF 文档中的文字（OCR）。把每一页渲染成图片后逐页识别，"
                    "适用于扫描件/图片型 PDF（文字型 PDF 请直接用 pdf_read）。"
                    "返回每页的文字内容。"
                ),
                parameters={
                    "type": "object",
                    "properties": {
                        "path": {
                            "type": "string",
                            "description": "PDF 文件路径（相对工作区）",
                        },
                        "max_pages": {
                            "type": "number",
                            "description": "最多识别页数（默认 20，从第 1 页开始）",
                        },
                    },
                    "required": ["path"],
                },
            ),
            ToolDefinition(
                name="ocr_pptx",
                description=(
                    "识别 PowerPoint (.pptx) 演示文稿中图片里的文字（OCR）。"
                    "逐页提取内嵌图片并识别，返回每张图片的文字与在页面的位置"
                    "（左上角坐标与宽高，按阅读顺序排序）。"
                ),
                parameters={
                    "type": "object",
                    "properties": {
                        "path": {
                            "type": "string",
                            "description": "pptx 文件路径（相对工作区）",
                        },
                    },
                    "required": ["path"],
                },
            ),
        ]

    # ------------------------------------------------------------ 执行入口

    async def execute(self, name: str, args: Dict[str, Any]) -> str:
        import asyncio as _asyncio

        handlers = {
            "ocr_image": self._ocr_image,
            "ocr_document": self._ocr_document,
            "ocr_pptx": self._ocr_pptx,
        }
        handler = handlers.get(name)
        if handler is None:
            raise ValueError(f"Unknown OCR Tool: {name}")
        return await _asyncio.to_thread(handler, args)

    def _resolve_path(self, rel_path: str) -> str:
        resolved = os.path.abspath(
            rel_path if os.path.isabs(rel_path) else os.path.join(self.workspace, rel_path)
        )
        if not (resolved == self.workspace or resolved.startswith(self.workspace + os.sep)):
            raise ValueError(f"路径越界：仅支持工作区内的文件: {rel_path}")
        if not os.path.isfile(resolved):
            raise ValueError(f"文件不存在: {rel_path}")
        return resolved

    # ------------------------------------------------------------ 图片

    def _ocr_image(self, args: Dict[str, Any]) -> str:
        if not _HAS_RAPIDOCR:
            return _missing_ocr_msg()
        rel_path = str(args.get("path") or "").strip()
        if not rel_path:
            return "[OCR Error]: 缺少 path 参数"
        try:
            resolved = self._resolve_path(rel_path)
        except ValueError as exc:
            return f"[OCR Error]: {exc}"
        ext = os.path.splitext(resolved)[1].lower()
        if ext not in IMAGE_EXTS:
            return f"[OCR Error]: 不支持的文件类型 {ext}（支持: {', '.join(sorted(IMAGE_EXTS))}）"
        try:
            text = _ocr_image_bytes(resolved)
        except Exception as exc:
            return f"[OCR Error]: 识别失败: {exc}"
        if not text.strip():
            return f"[OCR]: 未在图片中识别到文字（{os.path.basename(resolved)}）"
        return f"[OCR OK]: {os.path.basename(resolved)} 识别到 {len(text.splitlines())} 行\n\n{text[:MAX_TEXT_LEN]}"

    # ------------------------------------------------------------ PDF

    def _ocr_document(self, args: Dict[str, Any]) -> str:
        if not _HAS_RAPIDOCR:
            return _missing_ocr_msg()
        if not _HAS_PYMUPDF:
            return _missing_pdf_msg()
        rel_path = str(args.get("path") or "").strip()
        if not rel_path:
            return "[OCR Error]: 缺少 path 参数"
        try:
            resolved = self._resolve_path(rel_path)
        except ValueError as exc:
            return f"[OCR Error]: {exc}"
        if not resolved.lower().endswith(".pdf"):
            return "[OCR Error]: 仅支持 .pdf 文件"
        max_pages = max(1, min(100, int(args.get("max_pages") or 20)))

        try:
            doc = pymupdf.open(resolved)
        except Exception as exc:
            return f"[OCR Error]: 无法打开 PDF: {exc}"
        total = len(doc)
        parts = [f"PDF 共 {total} 页，识别前 {min(max_pages, total)} 页"]
        for i in range(min(max_pages, total)):
            try:
                pix = doc[i].get_pixmap(dpi=200)
                text = _ocr_pixmap(pix)
            except Exception as exc:
                text = f"[识别失败: {exc}]"
            parts.append(f"\n--- 第 {i + 1} 页 ---\n{text[:MAX_TEXT_LEN]}")
        doc.close()
        return "\n".join(parts)

    # ------------------------------------------------------------ PPTX（带位置）

    def _ocr_pptx(self, args: Dict[str, Any]) -> str:
        if not _HAS_RAPIDOCR:
            return _missing_ocr_msg()
        rel_path = str(args.get("path") or "").strip()
        if not rel_path:
            return "[OCR Error]: 缺少 path 参数"
        try:
            resolved = self._resolve_path(rel_path)
        except ValueError as exc:
            return f"[OCR Error]: {exc}"
        if not resolved.lower().endswith(".pptx"):
            return "[OCR Error]: 仅支持 .pptx 文件"
        try:
            from pptx import Presentation
        except ImportError:
            return "[OCR Error]: 需要 python-pptx，请运行: pip install python-pptx"

        import tempfile

        try:
            prs = Presentation(resolved)
        except Exception as exc:
            return f"[OCR Error]: 无法打开演示文稿: {exc}"

        parts = [f"共 {len(prs.slides)} 页幻灯片"]
        total_images = 0
        with tempfile.TemporaryDirectory(prefix="litework-ocr-") as tmpdir:
            for idx, slide in enumerate(prs.slides, 1):
                images = []
                for shape in slide.shapes:
                    if shape.shape_type is None:
                        continue
                    # 内嵌图片（Picture）
                    try:
                        st = str(getattr(shape, "shape_type", ""))
                        is_picture = "PICTURE" in st.upper() or getattr(shape, "image", None) is not None
                    except Exception:
                        is_picture = False
                    if not is_picture:
                        continue
                    try:
                        img = shape.image
                        ext_map = {"image/png": ".png", "image/jpeg": ".jpg",
                                   "image/gif": ".gif", "image/bmp": ".bmp"}
                        ext = ext_map.get(img.content_type, ".png")
                        tmp_path = os.path.join(tmpdir, f"s{idx}_{len(images)}{ext}")
                        with open(tmp_path, "wb") as f:
                            f.write(img.blob)
                        # 位置（EMU → 像素：1 px = 9525 EMU @96dpi）
                        left = getattr(shape, "left", 0) or 0
                        top = getattr(shape, "top", 0) or 0
                        width = getattr(shape, "width", 0) or 0
                        height = getattr(shape, "height", 0) or 0
                        images.append({
                            "path": tmp_path,
                            "x": round(left / 9525), "y": round(top / 9525),
                            "w": round(width / 9525), "h": round(height / 9525),
                        })
                    except Exception:
                        continue
                if not images:
                    continue
                total_images += len(images)
                parts.append(f"\n=== 第 {idx} 页（{len(images)} 张图片）===")
                # 按位置排序：先上后下（y），同一行先左后右（x）
                images.sort(key=lambda im: (im["y"], im["x"]))
                for im in images:
                    try:
                        text = _ocr_image_bytes(im["path"])
                    except Exception as exc:
                        text = f"[识别失败: {exc}]"
                    text = text.strip() or "（未识别到文字）"
                    parts.append(
                        f"\n【位置: 左上({im['x']},{im['y']}) 尺寸 {im['w']}×{im['h']}px】\n{text[:MAX_TEXT_LEN]}"
                    )
        if total_images == 0:
            return f"[OCR]: 演示文稿中没有找到内嵌图片（{len(prs.slides)} 页）"
        parts.insert(1, f"共识别 {total_images} 张图片中的文字")
        return "\n".join(parts)


def _ocr_pixmap(pix) -> str:
    """识别 pymupdf 渲染出的页面位图。"""
    import tempfile

    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
        tmp = f.name
    try:
        pix.save(tmp)
        return _ocr_image_bytes(tmp)
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass


# ---------------------------------------------------------------- 社区分发包装
# 由 litework 主仓库同步生成；安装到 ~/.lite-work/plugins/ 后覆盖内置同名
# 插件，卸载自动回退内置版。无参构造（插件加载器约定），workspace 在
# install 时从 kernel 的 app 服务捕获（项目热切换后新 kernel 重新 install）。

from litework.tools.plugin import ToolPlugin


class OcrPlugin(ToolPlugin):
    """ocr-plugin 社区独立分发版。"""

    name = "ocr-plugin"
    version = "1.0.0"
    description = "OCR 识别：图片/PDF 页面/PPT 内嵌图片中的文字提取"

    def __init__(self) -> None:
        self._app = None

    def install(self, kernel) -> None:
        try:
            if kernel.has_service("app"):
                self._app = kernel.get_service("app")
        except Exception:
            self._app = None
        super().install(kernel)

    def _workspace(self):
        return getattr(self._app, "workspace", None) if self._app else None

    def get_tools(self):
        return OCRTools(self._workspace()).get_tools()

    async def execute(self, name, args):
        return await OCRTools(self._workspace()).execute(name, args)
